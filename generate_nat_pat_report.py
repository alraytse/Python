#!/usr/bin/env python3
"""Generate a comprehensive NAT/PAT report from a collector CSV.

The input CSV is expected to contain these fields:
    Device, Platform, Record_Type, Command, Inside_Local, Inside_Global,
    Outside_Local, Outside_Global, Details

The script creates:
    1. A detailed CSV with NAT/PAT status and evidence fields.
    2. An Excel workbook with Summary and NAT_PAT_Detail worksheets.

Example:
    python generate_nat_pat_report.py \
        --input "switch_nat_report(in).csv" \
        --xlsx-output comprehensive_nat_pat_report.xlsx \
        --csv-output comprehensive_nat_pat_report.csv

Dependency:
    pip install openpyxl
"""

from __future__ import annotations

import argparse
import csv
import sys
from collections import OrderedDict
from pathlib import Path
from typing import Dict, Iterable, List

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    print("Missing dependency: openpyxl", file=sys.stderr)
    print("Install it with: pip install openpyxl", file=sys.stderr)
    raise SystemExit(2) from exc


DETAIL_FIELDS = [
    "Device",
    "Platform",
    "Record_Type",
    "NAT",
    "PAT",
    "Evidence_Type",
    "Protocol",
    "Command",
    "Inside_Local",
    "Inside_Global",
    "Outside_Local",
    "Outside_Global",
    "Details",
]

REQUIRED_FIELDS = {
    "Device",
    "Platform",
    "Record_Type",
    "Command",
    "Details",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate comprehensive NAT/PAT CSV and Excel reports."
    )
    parser.add_argument(
        "--input",
        default="switch_nat_report(in).csv",
        help="Input collector CSV. Default: switch_nat_report(in).csv",
    )
    parser.add_argument(
        "--xlsx-output",
        default="comprehensive_nat_pat_report.xlsx",
        help="Excel report output. Default: comprehensive_nat_pat_report.xlsx",
    )
    parser.add_argument(
        "--csv-output",
        default="comprehensive_nat_pat_report.csv",
        help="Detailed CSV report output. Default: comprehensive_nat_pat_report.csv",
    )
    parser.add_argument(
        "--encoding",
        default="cp1252",
        help="Input CSV encoding. Default: cp1252",
    )
    return parser.parse_args()


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def evidence_type(command: str) -> str:
    command_lower = command.lower()
    if "translation" in command_lower:
        return "Active Translation"
    if "statistic" in command_lower:
        return "PAT Statistics"
    if "running-config" in command_lower or "configuration" in command_lower:
        return "Configured Rule"
    return "Operational Output"


def protocol(details: str, record_type: str) -> str:
    if record_type != "NAT":
        return "N/A"

    first_token = details.split(None, 1)[0].lower() if details.split() else ""
    if first_token in {"tcp", "udp", "icmp", "gre", "ip"}:
        return first_token.upper()
    return "N/A"


def read_source_csv(path: Path, encoding: str) -> List[Dict[str, str]]:
    with path.open("r", encoding=encoding, newline="") as source_file:
        reader = csv.DictReader(source_file)
        fieldnames = {clean(field) for field in (reader.fieldnames or [])}
        missing = REQUIRED_FIELDS - fieldnames
        if missing:
            missing_text = ", ".join(sorted(missing))
            raise ValueError(f"Input CSV is missing required columns: {missing_text}")

        rows: List[Dict[str, str]] = []
        for source_row in reader:
            row = {clean(key): clean(value) for key, value in source_row.items() if key}
            rows.append(row)
        return rows


def enrich_rows(source_rows: Iterable[Dict[str, str]]) -> List[Dict[str, str]]:
    enriched: List[Dict[str, str]] = []
    for source_row in source_rows:
        record_type = clean(source_row.get("Record_Type")).upper() or "N/A"
        details = clean(source_row.get("Details"))
        command = clean(source_row.get("Command"))

        enriched.append(
            {
                "Device": clean(source_row.get("Device")),
                "Platform": clean(source_row.get("Platform")),
                "Record_Type": record_type,
                "NAT": "Configured" if record_type == "NAT" else "N/A",
                "PAT": "Configured" if record_type == "PAT" else "N/A",
                "Evidence_Type": evidence_type(command),
                "Protocol": protocol(details, record_type),
                "Command": command,
                "Inside_Local": clean(source_row.get("Inside_Local")),
                "Inside_Global": clean(source_row.get("Inside_Global")),
                "Outside_Local": clean(source_row.get("Outside_Local")),
                "Outside_Global": clean(source_row.get("Outside_Global")),
                "Details": details,
            }
        )
    return enriched


def ordered_devices(rows: Iterable[Dict[str, str]]) -> List[str]:
    devices = OrderedDict()
    for row in rows:
        devices.setdefault(row["Device"], None)
    return list(devices)


def write_detail_csv(rows: List[Dict[str, str]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8-sig", newline="") as output_file:
        writer = csv.DictWriter(output_file, fieldnames=DETAIL_FIELDS)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def style_title(sheet, title: str, end_column: int) -> None:
    navy = "17365D"
    white = "FFFFFF"
    sheet.cell(1, 1, title)
    sheet.cell(1, 1).font = Font(name="Arial", size=16, bold=True, color=white)
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor=navy)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    sheet.row_dimensions[1].height = 26


def style_header_row(sheet, row_number: int, headers: List[str]) -> None:
    navy = "17365D"
    white = "FFFFFF"
    thin_gray = Side(style="thin", color="B7C9D6")

    for column, header in enumerate(headers, 1):
        cell = sheet.cell(row_number, column, header)
        cell.font = Font(name="Arial", bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_gray)


def build_workbook(
    rows: List[Dict[str, str]],
    input_path: Path,
    output_path: Path,
) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    detail = wb.create_sheet("NAT_PAT_Detail")

    for sheet in (summary, detail):
        sheet.sheet_view.showGridLines = False

    style_title(summary, "Comprehensive NAT and PAT Report", 13)
    summary["A2"] = "Source file"
    summary["B2"] = input_path.name
    summary["A3"] = "Scope note"
    summary["B3"] = (
        "The source contains operational NAT translations and PAT statistics. "
        "Status is based on record types present in the supplied CSV."
    )
    summary.merge_cells("B3:M3")
    summary["A4"] = "Status rule"
    summary["B4"] = (
        "NAT or PAT is Configured when at least one corresponding record exists; "
        "otherwise N/A."
    )
    summary.merge_cells("B4:M4")

    for cell_name in ("A2", "A3", "A4"):
        summary[cell_name].font = Font(name="Arial", bold=True, color="17365D")
    for cell_name in ("B2", "B3", "B4"):
        summary[cell_name].font = Font(name="Arial", italic=cell_name != "B2")

    summary_headers = [
        "Device",
        "Platform",
        "NAT",
        "PAT",
        "NAT_Record_Count",
        "PAT_Record_Count",
        "Total_Record_Count",
        "Active_Translation_Count",
        "PAT_Statistics_Count",
        "NAT_Commands",
        "PAT_Commands",
        "PAT_Details",
        "Notes",
    ]
    summary_header_row = 6
    style_header_row(summary, summary_header_row, summary_headers)

    detail_start_row = 5
    detail_end_row = detail_start_row + len(rows) - 1
    devices = ordered_devices(rows)
    thin_gray = Side(style="thin", color="B7C9D6")
    green = "E2F0D9"

    for output_row, device in enumerate(devices, summary_header_row + 1):
        device_rows = [row for row in rows if row["Device"] == device]
        platform = device_rows[0]["Platform"] if device_rows else "N/A"
        nat_commands = sorted(
            {row["Command"] for row in device_rows if row["Record_Type"] == "NAT" and row["Command"]}
        )
        pat_commands = sorted(
            {row["Command"] for row in device_rows if row["Record_Type"] == "PAT" and row["Command"]}
        )
        pat_details = sorted(
            {row["Details"] for row in device_rows if row["Record_Type"] == "PAT" and row["Details"]}
        )

        summary.cell(output_row, 1, device)
        summary.cell(output_row, 2, platform)
        summary.cell(output_row, 3, f'=IF(E{output_row}>0,"Configured","N/A")')
        summary.cell(output_row, 4, f'=IF(F{output_row}>0,"Configured","N/A")')
        summary.cell(
            output_row,
            5,
            f'=COUNTIFS(NAT_PAT_Detail!$A${detail_start_row}:$A${detail_end_row},A{output_row},'
            f'NAT_PAT_Detail!$C${detail_start_row}:$C${detail_end_row},"NAT")',
        )
        summary.cell(
            output_row,
            6,
            f'=COUNTIFS(NAT_PAT_Detail!$A${detail_start_row}:$A${detail_end_row},A{output_row},'
            f'NAT_PAT_Detail!$C${detail_start_row}:$C${detail_end_row},"PAT")',
        )
        summary.cell(output_row, 7, f"=SUM(E{output_row}:F{output_row})")
        summary.cell(
            output_row,
            8,
            f'=COUNTIFS(NAT_PAT_Detail!$A${detail_start_row}:$A${detail_end_row},A{output_row},'
            f'NAT_PAT_Detail!$F${detail_start_row}:$F${detail_end_row},"Active Translation")',
        )
        summary.cell(
            output_row,
            9,
            f'=COUNTIFS(NAT_PAT_Detail!$A${detail_start_row}:$A${detail_end_row},A{output_row},'
            f'NAT_PAT_Detail!$F${detail_start_row}:$F${detail_end_row},"PAT Statistics")',
        )
        summary.cell(output_row, 10, ", ".join(nat_commands) or "N/A")
        summary.cell(output_row, 11, ", ".join(pat_commands) or "N/A")
        summary.cell(output_row, 12, "; ".join(pat_details) or "N/A")
        summary.cell(output_row, 13, "Review configuration if operational evidence is unexpected")

        for column in range(1, 14):
            cell = summary.cell(output_row, column)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
        summary.cell(output_row, 3).fill = PatternFill("solid", fgColor=green)
        summary.cell(output_row, 4).fill = PatternFill("solid", fgColor=green)

    style_title(detail, "NAT/PAT Detailed Records", len(DETAIL_FIELDS))
    detail["A2"] = "N/A indicates that the opposite rule type is not represented by that record."
    detail["A2"].font = Font(name="Arial", italic=True)
    detail.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(DETAIL_FIELDS))
    style_header_row(detail, 4, DETAIL_FIELDS)

    blue = "D9EAF7"
    gray = "E7E6E6"
    for output_row, row in enumerate(rows, detail_start_row):
        for column, field in enumerate(DETAIL_FIELDS, 1):
            value = row[field] or "N/A"
            cell = detail.cell(output_row, column, value)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
        detail.cell(output_row, 4).fill = PatternFill(
            "solid", fgColor=green if row["NAT"] == "Configured" else gray
        )
        detail.cell(output_row, 5).fill = PatternFill(
            "solid", fgColor=blue if row["PAT"] == "Configured" else gray
        )

    summary.freeze_panes = "A7"
    detail.freeze_panes = "A5"

    summary_widths = [24, 16, 14, 14, 16, 16, 17, 22, 20, 32, 32, 42, 42]
    detail_widths = [24, 16, 14, 14, 14, 20, 12, 28, 18, 18, 18, 18, 48]
    for column, width in enumerate(summary_widths, 1):
        summary.column_dimensions[get_column_letter(column)].width = width
    for column, width in enumerate(detail_widths, 1):
        detail.column_dimensions[get_column_letter(column)].width = width

    for sheet in (summary, detail):
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def validate_workbook(path: Path) -> None:
    workbook = load_workbook(path, data_only=False)
    if workbook.sheetnames != ["Summary", "NAT_PAT_Detail"]:
        raise ValueError("Generated workbook does not contain the expected worksheets")
    if workbook["Summary"]["C7"].value is None:
        raise ValueError("Generated workbook is missing summary status formulas")
    if workbook["NAT_PAT_Detail"]["D5"].value not in {"Configured", "N/A"}:
        raise ValueError("Generated workbook is missing NAT status values")


def main() -> int:
    args = parse_args()
    input_path = Path(args.input).expanduser().resolve()
    xlsx_path = Path(args.xlsx_output).expanduser().resolve()
    csv_path = Path(args.csv_output).expanduser().resolve()

    if not input_path.is_file():
        print(f"Input file not found: {input_path}", file=sys.stderr)
        return 1

    try:
        source_rows = read_source_csv(input_path, args.encoding)
        report_rows = enrich_rows(source_rows)
        write_detail_csv(report_rows, csv_path)
        build_workbook(report_rows, input_path, xlsx_path)
        validate_workbook(xlsx_path)
    except (OSError, ValueError, csv.Error) as exc:
        print(f"Report generation failed: {exc}", file=sys.stderr)
        return 1

    print(f"Input records: {len(source_rows)}")
    print(f"Devices: {len(ordered_devices(report_rows))}")
    print(f"Detailed CSV: {csv_path}")
    print(f"Excel report: {xlsx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
